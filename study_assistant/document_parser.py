from __future__ import annotations

from io import BytesIO
from pathlib import Path


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from supported learning-document formats."""
    suffix = Path(filename).suffix.lower()
    if suffix in {".txt", ".md"}:
        return content.decode("utf-8", errors="ignore")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(content)).pages)
    if suffix == ".docx":
        from docx import Document

        document = Document(BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(BytesIO(content))
        return "\n".join(
            shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if hasattr(shape, "text")
        )
    raise ValueError("Unsupported file type.")


SUPPORTED_EXTENSIONS = (".txt", ".md", ".pdf", ".docx", ".pptx")
